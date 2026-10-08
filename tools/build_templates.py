#!/usr/bin/env python3
"""
DiploClear — template builder.
Turns each host nation's REAL form into a docxtemplater template by replacing
only the sample VALUES with {canonical_placeholders}. Layout, labels, tables and
classification banners stay exactly as the host wrote them (format-in = format-out).
Single-leg for now; multi-leg looping is a later refinement.
Re-run:  python3 tools/build_templates.py
"""
import os, base64, re, docx
from docx.oxml.ns import qn

SRC = "/Users/xynkro/Documents/1. CSC/Technopreneurial Mindset (TM)/TM Project/TM (Dip Forms)"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOFFICE = "/Users/xynkro/.claude/skills/pptx/scripts/office/soffice.py"

def scrub_docprops(path):
    """Clear the Word/Excel authorship metadata.

    The scrub has always covered body text, but every .docx also carries docProps: the
    creator, the last person to save it, their company. Twelve of these templates were
    stamped "Ronnie, SO2 DPCS, AOCG" as last-modified-by, and that travelled with the file
    to twelve host nations on every request. Nobody reads it, which is exactly why it
    survived. Returns the fields it cleared."""
    import zipfile, shutil, tempfile, re as _re
    FIELDS = ["dc:creator", "cp:lastModifiedBy", "Company", "Manager", "dc:description",
              "cp:category", "cp:keywords", "dc:subject", "cp:lastPrinted"]
    cleared = []
    zin = zipfile.ZipFile(path)
    items = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    for name in ("docProps/core.xml", "docProps/app.xml"):
        if name not in items: continue
        x = items[name].decode("utf-8", "replace")
        for tag in FIELDS:
            pat = _re.compile(r"<%s(\s[^>]*)?>([^<]*)</%s>" % (_re.escape(tag), _re.escape(tag)))
            m = pat.search(x)
            if m and m.group(2).strip():
                cleared.append(f"{tag.split(':')[-1]}={m.group(2).strip()[:28]}")
                x = pat.sub(lambda mm: f"<{tag}{mm.group(1) or ''}></{tag}>", x)
        items[name] = x.encode("utf-8")
    tmp = tempfile.mktemp(suffix=".zip")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in items.items(): zout.writestr(n, data)
    shutil.move(tmp, path)
    return cleared


def stamp_cache_busters():
    """Rewrite the ?v= on each <script src="lib/..."> to that file's content hash.
    A planner who already has the app open holds a cached templates_data.js; without a
    changing URL a rebuilt template never reaches them and they keep filing the old one."""
    import hashlib
    idx_path = os.path.join(ROOT, "index.html")
    html = open(idx_path, encoding="utf-8").read()
    changed = []
    for lib in ("pizzip.js", "docxtemplater.js", "pdf-lib.js", "templates_data.js"):
        fp = os.path.join(ROOT, "lib", lib)
        if not os.path.exists(fp): continue
        h = hashlib.sha1(open(fp, "rb").read()).hexdigest()[:10]
        new = f'<script src="lib/{lib}?v={h}"></script>'
        pat = re.compile(r'<script src="lib/' + re.escape(lib) + r'(?:\?v=[0-9a-f]+)?"></script>')
        if not pat.search(html):
            print(f"  !! cache-buster: no script tag for {lib}"); continue
        if pat.search(html).group(0) != new: changed.append(lib)
        html = pat.sub(new, html, count=1)
    open(idx_path, "w", encoding="utf-8").write(html)
    print(f"  cache-busters stamped{' (updated: ' + ', '.join(changed) + ')' if changed else ' (unchanged)'}")


def repl(p, old, new, misses, tag):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new); return
    full = "".join(r.text for r in p.runs)
    if old in full:
        for r in p.runs: r.text = ""
        if p.runs: p.runs[0].text = full.replace(old, new)
        return
    misses.append(f"{tag}:'{old[:34]}'")

def setp(p, text):
    for r in p.runs: r.text = ""
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)

def clearp(p):
    for r in p.runs: r.text = ""

# ---- table-cell helpers (for EU DIC forms, where data lives in a table) ----
def set_cell(cell, text):
    for p in cell.paragraphs:
        for r in p.runs: r.text = ""
    p0 = cell.paragraphs[0]
    if p0.runs: p0.runs[0].text = text
    else: p0.add_run(text)

def tset(doc, t, marker, new):
    for row in doc.tables[t].rows:
        for cell in row.cells:
            if marker in cell.text:
                set_cell(cell, new)

def trepl(doc, t, old, new, misses):
    found = False
    for row in doc.tables[t].rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full = "".join(r.text for r in p.runs)
                if old not in full: continue
                found = True; done = False
                for r in p.runs:
                    if old in r.text: r.text = r.text.replace(old, new); done = True
                if not done:
                    for r in p.runs: r.text = ""
                    if p.runs: p.runs[0].text = full.replace(old, new)
    if not found: misses.append(f"T{t}:'{old[:26]}'")

def tform(doc, t, label, new, misses):
    """Label-driven: set the LAST cell of the row whose label cell contains `label`."""
    hit = False
    for row in doc.tables[t].rows:
        if any(label in c.text for c in row.cells):
            set_cell(row.cells[-1], new); hit = True
    if not hit: misses.append(f"T{t}/{label[:18]}")

def tcol(doc, t, label, col, new, misses):
    """Like tform but sets a SPECIFIC column of the matched row (for Oman: value lives in col 2)."""
    hit = False
    for row in doc.tables[t].rows:
        cells = row.cells
        if any(label in c.text for c in cells) and col < len(cells):
            set_cell(cells[col], new); hit = True
    if not hit: misses.append(f"T{t}/{label[:16]}")

def del_tables(doc, indices):
    tbls = list(doc.tables)
    for i in indices:
        tbls[i]._element.getparent().remove(tbls[i]._element)

def cut_after(doc, marker):
    body = doc.element.body
    target = None
    for child in body:
        if child.tag == qn('w:p'):
            if marker in "".join(n.text or "" for n in child.iter(qn('w:t'))):
                target = child; break
    if target is None: return 0
    els, cur = [], target
    while cur is not None:
        els.append(cur); cur = cur.getnext()
    n = 0
    for el in els:
        if el.tag == qn('w:sectPr'): continue
        body.remove(el); n += 1
    return n

def resolve_src(name):
    p = os.path.join(SRC, name)
    if name.lower().endswith('.doc'):
        outdir = '/tmp/_conv'; os.makedirs(outdir, exist_ok=True)
        os.system(f'python3 "{SOFFICE}" --headless --convert-to docx --outdir "{outdir}" "{p}" >/dev/null 2>&1')
        return os.path.join(outdir, os.path.splitext(os.path.basename(name))[0] + '.docx')
    return p

def build(cfg):
    d = docx.Document(resolve_src(cfg["src"]))
    P = d.paragraphs; misses = []
    for idx, old, new in cfg.get("repls", []):
        repl(P[idx], old, new, misses, f"P{idx}")
    for idx, text in cfg.get("setps", []):
        setp(P[idx], text)
    for idx in cfg.get("clears", []):
        clearp(P[idx])
    for t, old, new in cfg.get("tcells", []):
        trepl(d, t, old, new, misses)
    for t, label, new in cfg.get("tform", []):
        tform(d, t, label, new, misses)
    for t, label, col, new in cfg.get("tcol", []):
        tcol(d, t, label, col, new, misses)
    for t, marker, new in cfg.get("tset", []):
        tset(d, t, marker, new)
    if cfg.get("del_tables"):
        del_tables(d, cfg["del_tables"])
    cut = cut_after(d, cfg["cut_marker"]) if cfg.get("cut_marker") else 0
    outdir = os.path.join(ROOT, "Master Forms", cfg["country"]); os.makedirs(outdir, exist_ok=True)
    outpath = os.path.join(outdir, cfg["key"] + ".docx"); d.save(outpath)
    meta = scrub_docprops(outpath)
    doc2 = docx.Document(outpath)
    txt = "\n".join(p.text for p in doc2.paragraphs)
    for tb in doc2.tables:
        for row in tb.rows:
            for cell in row.cells: txt += "\n" + cell.text
    tags = sorted(set(re.findall(r"\{(\w+)\}", txt)))
    leftover = [s for s in cfg.get("sentinels", []) if s in txt]
    flag = "  ⚠" if (leftover or misses) else ""
    print(f"  {cfg['country']:12} legs_cut={cut:<3} tags={len(tags):<2} leftover={leftover or '-'}"
          f"{' meta_cleared=' + str(len(meta)) if meta else ''}{(' MISS='+str(misses)) if misses else ''}{flag}")
    return cfg["key"], outpath

# Shared EU DIC form (Greece/Spain/Italy share one template; the data lives in TABLE 2,
# only the country name + routing block differ). Values are found by search, so row-offset-proof.
DIC_FORM = [(2,"Number and type of aircraft","{aircraft}"),(2,"Aircraft registration","{reg}"),
            (2,"sign","{callsign}"),(2,"Number of crew","{crew_n}"),(2,"Pilot rank and name","{captain}"),
            (2,"Date of flight","{date}"),(2,"Purpose of flight","{purpose}"),
            (2,"Departure airport","{dep}"),(2,"Destination airport","{dest}"),(2,"Number of passengers","{pax}")]

# ---------------- per-form configs (indices from structure dumps) ----------------
CONFIGS = [
 dict(country="Thailand", key="thailand", src="THAILAND_sanitized.docx", cut_marker="DCR 042-26 (2)",
   repls=[(0,"042-26 (1)","{dcr_no}"),(0,"06 MAR 26","{dtd}"),(10,"1 x A330 MRTT","{aircraft}"),
     (11,"SINGA 8613","{callsign}"),(12,"761 (ALT: 760, 762, 763, 764, 765 Or As Per Flt Plan)","{tails}"),
     (14,"Air to Air Refueling Flight -EFTD (Thai) 26-2","{purpose}"),(15,"MAJ Michelle Teo","{captain}"),
     (16,"10","{crew_n}"),(16,"Nil","{pax}"),(17,"FAK & Baggage","{cargo}"),(18,"Nil","{mil_equip}")],
   tcells=[
           # itinerary times: rows 1-3 from the entered legs, return rows cleared
           (0,"08 / 2300Z Apr 26","{p1_etd}"),(0,"09 / 0001Z Apr 26","{p2_eta}"),
           (0,"09 / 0200Z Apr 26","{p3_eta}"),(0,"09 / 0400Z Apr 26","{p3_etd}"),
           (0,"09 / 0600Z Apr 26",""),(0,"09 / 0700Z Apr 26",""),],
   sentinels=["SINGA 8613","Michelle Teo","2300Z","0001Z","0200Z","0400Z","0600Z","0700Z","Apr 26"]),

 dict(country="Vietnam", key="vietnam", src="VIETNAM_sanitized.doc", cut_marker="117-15(2)",
   repls=[(0,"117-15(1)","{dcr_no}"),(1,"03 AUG 15","{dtd}"),(6,"02 SEP 15","{date}"),(8,"1 x F-50","{aircraft}"),
     (10,"712 (ALT: 710, 711, 713)","{tails}"),(12,"SINGA 5549","{callsign}"),
     (14,"AIRLIFT  FLIGHT - (OVER- FLIGHT)","{purpose}"),(16,"(ALL TIMES UTC SEP 15)","{time_note}"),
     (39,"LTC CLEMENT WEE","{captain}"),(39,"07","{crew_n}"),(40,"STORES & BAGGAGES","{cargo}"),(44,"NIL","{photo}")],
   setps=[(24,"{sector}     :     {route}"),(31,"AS PER FILED FLIGHT PLAN")],
   clears=[18,19,20,21,26,27,32,33,35,36,37],
   sentinels=["SINGA 5549","CLEMENT WEE","SIEM REAP","SEP 15","HO CHI MINH"]),

 dict(country="Malaysia", key="malaysia", src="(MALAYSIA)_sanitized.docx", cut_marker="DCR 040-26 (2)",
   repls=[(1,"040-26 (1)","{dcr_no}"),(2,"09 MAR 26","{dtd}"),(6,"08 APR 26","{date}"),(8,"1 x C130","{aircraft}"),
     (10,"730 (ALT: 720, 721, 724, 725, 731, 732,","{tails}"),(12,"SINGA 9635","{callsign}"),
     (14,"IN SUPPORT OF RSAF THAILAND","{purpose}"),(17,"PAYA LEBAR - UDON THANI","{sector}"),
     (23,"LTA JOSHUA LIM WEE CHONG","{captain}"),(25,"STORES / BAGGAGE / EXPLOSIVE CL:","{cargo}"),
     (30,"NIL","{photo}"),(34,"NIL","{services}"),(36,"72 HOURS.","{validity}")],
   setps=[(20,"{sector}"),(21,"{route}")], clears=[11,15,26],
   tcells=[
           # itinerary times from the entered legs
           (0,"07 APR - 2300 UTC","{p1_etd}"),(0,"07 APR - 2305 UTC","{p2_eta}"),
           (0,"07 APR - 2335 UTC","{p2_etd}"),(0,"08 APR - 0300 UTC","{p3_eta}"),],
   sentinels=["SINGA 9635","JOSHUA LIM","2300 UTC","2305 UTC","2335 UTC","0300 UTC","07 APR"]),

 dict(country="Cambodia", key="cambodia", src="CAMBODIA_sanitized.doc", cut_marker="DCR 117-15 (2)",
   repls=[(0,"117-15 (1)","{dcr_no}"),(1,"03 AUG 15","{dtd}"),(5,"1 x F-50","{aircraft}"),
     (7,"712 (ALT 711, 710, 713)","{tails}"),(9,"SINGA 5549 / 02 SEP 15","{callsign} / {date}"),
     (11,"CHANGI - SIEM REAP -","{sector}"),(14,"(ALL TIMES UTC SEP 15)","{time_note}"),
     (27,"FL 200 - 290","{level}"),(29,"VIP AIRLIFT FLIGHT IN SUPPORT OF 12TH","{purpose}"),
     (32,"LTC CLEMENT WEE","{captain}"),(32,"07","{crew_n}"),(34,"TBC","{pax}"),
     (36,"BAGGAGE / FAK","{cargo}"),(38,"CHANGI","{dep}"),(40,"NIL","{photo}")],
   setps=[(22,"{sector}     :"),(23,"{route}")], clears=[12,16,17,18,19,24,25,30],
   sentinels=["SINGA 5549","CLEMENT WEE","SIEM REAP","SEP 15"]),

 dict(country="Myanmar", key="myanmar", src="MYANMAR_sanitized.doc", cut_marker="DCR 068-25 (2)",
   repls=[(0,"068-25 (1)","{dcr_no}"),(1,"02 APR 25","{dtd}"),(5,"1 x C-130","{aircraft}"),
     (7,"730 (ALT: 720, 721, 724, 725, 731, 732,","{tails}"),(10,"SINGA 9560  (04 APR 25)","{callsign}  ({date})"),
     (12,"PAYA LEBAR - YANGON INT’L - PAYA LEBAR","{sector}"),(14,"(ALL TIMES UTC - APR 25)","{time_note}"),
     (32,"FL 180 - 250","{level}"),(34,"HUMANITARIAN & DISASTER RELIEF SUPPORT","{purpose}"),
     (37,"MAJ C TEENESHWARAN","{captain}"),(37,"20","{crew_n}"),(39,"11","{pax}"),
     (41,"HADR STORES","{cargo}"),(43,"PAYA LEBAR AIR BASE","{dep}"),(45,"NIL","{photo}")],
   setps=[(23,"AS PER FILED FLIGHT PLAN"),(27,"{sector}"),(28,"{route}")],
   clears=[8,16,17,18,19,24,29,30,35],
   sentinels=["SINGA 9560","TEENESHWARAN","YANGON","APR 25"]),

 dict(country="Saudi Arabia", key="saudi", src="SAUDI ARABIA_sanitized.doc", cut_marker="DCR 060-26 (2)",
   repls=[(0,"060-26 (1)","{dcr_no}"),(1,"06 MAR 26","{dtd}"),(5,"1 x A330 MRTT","{aircraft}"),
     (7,"764 (ALT: 760, 761, 762, 763, 765 Or As","{tails}"),(10,"SINGA 12 / 10 MAR 26","{callsign} / {date}"),
     (12,"CHANGI AIR BASE, SINGAPORE -","{sector}"),(16,"(ALL TIMES UTC)","{time_note}"),
     (30,"FL 350 - 390","{level}"),(32,"REPATRIATION FLIGHT","{purpose}"),
     (34,"LTC LEE TAT WEE","{captain}"),(34,"20","{crew_n}"),(36,"NIL","{pax}"),
     (38,"STORES & PERSONAL BAGGAGE","{cargo}"),(40,"CHANGI AIR BASE (WSSS)","{dep}"),(42,"NIL","{photo}")],
   setps=[(24,"AS PER FILED FLIGHT PLAN"),(28,"{route}")], clears=[8,13,14,18,19,20,25],
   sentinels=["LEE TAT WEE","SITOL","0630"]),

 dict(country="Taiwan", key="taiwan", src="TAIWAN_sanitized.doc", cut_marker="DCR 034-20 (2)",
   repls=[(0,"034-20 (1)","{dcr_no}"),(1,"10 JAN 20","{dtd}"),(5,"1 x A330 MRTT","{aircraft}"),
     (7,"763 (ALT: 760, 761, 762 or As per flt plan)","{tails}"),(9,"SINGA 82 (16 JAN 20)","{callsign} ({date})"),
     (11,"CHANGI - KAOHSIUNG INT’L - CHANGI","{sector}"),(13,"(ALL TIMES UTC - JAN 20)","{time_note}"),
     (29,"AEROMEDICAL EVACUATION FLIGHT","{purpose}"),(31,"MAJ CHAN WEE WEE","{captain}"),(31,"03","{crew_n}"),
     (33,"32","{pax}"),(35,"AEROMEDICAL PATIENT & BAGGAGE","{cargo}"),(39,"NIL","{photo}")],
   setps=[(21,"{sector}     :     {route}")], clears=[15,16,17,18,22,23,24,26,27],
   sentinels=["SINGA 82","CHAN WEE WEE","KAOHSIUNG","JAN 20"]),

 dict(country="Greece", key="greece", src="GREECE_sanitized.docx", tform=DIC_FORM,
   tset=[(4,"ENTER:","AS PER FILED FLIGHT PLAN"),(4,"EXIT:",""),(4,"BELIX","{route}")],
   sentinels=["SINGA 23","CRUZ LOPEZ","RSAF FERRY","BELIX","BARAJAS"]),

 dict(country="Spain", key="spain", src="SPAIN_sanitized.docx", tform=DIC_FORM,
   tset=[(4,"ETD:","AS PER FILED FLIGHT PLAN"),(4,"EXIT:",""),(4,"PINAR","{route}")],
   sentinels=["SINGA 23","CRUZ LOPEZ","RSAF FERRY","PINAR","BARAJAS"]),

 dict(country="Italy", key="italy", src="ITALY_sanitized.doc", tform=DIC_FORM,
   tset=[(4,"ENTER:","AS PER FILED FLIGHT PLAN"),(4,"EXIT:",""),(4,"ELSAG","{route}")],
   sentinels=["SINGA 23","CRUZ LOPEZ","RSAF FERRY","ELSAG","BARAJAS"]),

 # ---- bespoke Word forms ----
 dict(country="Germany", key="germany", src="GERMANY TEMPLATE_sanitized.docx",
   tcells=[(0,"16 MAR 20","{date}"),(1,"CHANGI AIR BASE (WSSS)","{dep}"),(1,"DRESDEN AIRPORT (EDDC)","{dest}"),
           (2,"A330 MRTT","{type}"),(2,"SINGA 80","{callsign}"),(2,"24 MAR 20","{date}"),(2,"DRESDEN AIRPORT","{dest}"),
           (5,"REPATRIATION FLIGHT","{purpose}")],
   tset=[(4,"WSSS","{pax}")],
   sentinels=["SINGA 80","DRESDEN","16 MAR 20"]),

 # Egypt prints airspace entry/exit with their own times. Those were the sample ferry
 # flight's (SALUN 1405Z, IMRAD 1555Z on 12 MAR 26) and shipped on every request.
 dict(country="Egypt", key="egypt", src="EGYPT_sanitized.docx",
   repls=[(9,"760","{reg}"),(10,"12 MAR 2026","{date}"),(12,"A330 MRTT","{type}"),(13,"SINGA 23","{callsign}"),
          (19,"CHANGI AIRBASE","{dest}"),(20,"BARAJAS AIRPORT","{dep}"),(32,"RSAF FERRY FLIGHT","{purpose}"),
          (18,"SALUN","{fir_entry}"),(18,"1405Z","{fir_entry_t}"),(18,"12 MAR 2026","{date}"),
          (19,"0130Z","{eta}"),(19,"13 MAR 2026","{date}"),
          (20,"1100Z","{etd}"),(20,"12 MAR 2026","{date}"),
          (21,"IMRAD","{fir_exit}"),(21,"1555Z","{fir_exit_t}"),(21,"12 MAR 2026","{date}")],
   sentinels=["SINGA 23","BARAJAS","RSAF FERRY","SALUN","IMRAD","1405Z","1555Z","0130Z","1100Z","MAR 2026"]),

 dict(country="Oman", key="oman", src="OMAN_sanitized_sanitized.docx", del_tables=[1,2,3],
   tcol=[(0,"Aircraft Type",2,"{type}"),(0,"Registration:",2,"{tails}"),(0,"Call sign",2,"{callsign}"),
         (0,"Date of operation",2,"{date}"),(0,"Origin (airport",2,"{dep}"),(0,"Destination (airport",2,"{dest}"),
         (0,"Mission:",2,"{purpose}"),(0,"Lead passenger",2,"{captain}"),(0,"Number of passengers",2,"{pax}"),
         (0,"Nature of Cargo",2,"{cargo}")],
   sentinels=["SINGA 14","CHEN JIANWEI","ABDULAZIZ"]),

 # India = multi-table Indian "Application for Non-Scheduled Flights" (APPLICATION FORM paragraphs + tables)
 # India's form is a round trip: T0/T5-row-1 = outbound, T2/T5-row-2 = return. Until an
 # inbound leg is entered the return side is BLANKED rather than left carrying the sample
 # mission's Riyadh routing — the form used to tell India we were flying OERK-WSSS on 10 MAR
 # whatever the real sortie was. A blank the planner completes beats a confident wrong answer.
 dict(country="India", key="india", src="DCR 060-26 (INDIA).doc (SINGA 12 )_sanitized.doc",
   repls=[(68,"A330 MRTT","{type}"),(69,"SINGA 12","{callsign}"),
          (70,"764 (ALT : 760, 761, 762, 763, 765 Or As per flt plan)","{tails}"),
          (79,"REPATRIATION FLIGHT.","{purpose}"),
          (87,"HADR STORES, BAGGAGE & FLY AWAY KIT","{cargo}")],
   tcells=[(0,"Repatriation Flight","{purpose}"),(0,"A330 MRTT","{type}"),(0,"LTC LEE TAT WEE","{captain}"),
           (2,"Repatriation Flight","{purpose}"),(2,"A330 MRTT","{type}"),(2,"LTC LEE TAT WEE","{captain}"),
           (5,"09 MAR 25","{date}"),
           # --- outbound: real itinerary ---
           (0,"CHANGI \u2013 KING KHALED INT\u2019L","{sector}"),
           (0,"WSSS DCT SJ DCT SALAX N563 REXOD L883 KITUB Y517 TOTEB DCT OERK","{route}"),
           (0,"Entry Point: MEMAK @ 0030Z","Entry Point: {fir_entry} @ {fir_entry_t}"),
           (0,"Exit Point:    KITAL    @ 0515Z","Exit Point: {fir_exit} @ {fir_exit_t}"),
           (0,"ETD: WSSS - 09 MAR 26 @ 2300Z","ETD: {dep} - {etd_full}"),
           (0,"ETA: OERK  - 10 MAR 26 @ 0800Z","ETA: {dest} - {eta_full}"),
           (1,"Date: 06 MAR 26","Date: {date}"),
           (3,"Date: 06 MAR 26","Date: {date}"),
           # --- return leg: cleared, not inherited ---
           (2,"KING KHALED INT\u2019L AIRPORT - CHANGI",""),
           (2,"OERK DCT TOTEB DCT NAGBU Y214 RAPMA DCT DAPOL L692 GISKA DCT UMILA L883 REXOD N563 SALAX A576 SJ DCT WSSS",""),
           (2,"Entry Point: REXOD @  1245Z","Entry Point:"),
           (2,"Exit Point:   MEMAK @  1730Z","Exit Point:"),
           (2,"ETD: OERK  - 10 MAR 26 @  1100Z","ETD:"),
           (2,"ETA: WSSS  - 10 MAR 26 @  1900Z","ETA:"),
           # --- summary table: row 1 outbound, row 2 return ---
           (5,"2300Z","{etd}"),(5,"0800Z","{eta}"),
           (5,"Entry Point: MEMAK @ 0030Z","Entry Point: {fir_entry} @ {fir_entry_t}"),
           (5,"REXOD @ 0515Z","{fir_exit} @ {fir_exit_t}"),
           (5,"10 MAR 26",""),(5,"1100Z",""),(5,"1900Z",""),
           (5,"Entry Point: REXOD @ 1245Z","Entry Point:"),(5,"MEMAK @ 1730Z",""),
           (5,"OERK DCT TOTEB DCT NAGBU Y214 RAPMA DCT DAPOL L692 GISKA DCT UMILA L883 REXOD N563 SALAX A576  SJ DCT WSSS",""),
           (5,"OERK","")],
   sentinels=["HADR STORES","LEE TAT WEE","Repatriation Flight",
              "MEMAK","KITAL","REXOD","OERK","0030Z","0515Z","1245Z","1730Z","MAR 26","MAR 25"]),
]

def build_france():
    """France = the EU DIC form in .xlsx. Set the data cells to {placeholders};
    the bilingual label formulas + all styling stay untouched. Runtime fills via
    text-replace in xl/sharedStrings.xml (same bundle-and-replace trick as docx)."""
    import openpyxl, warnings
    warnings.simplefilter("ignore")
    src = os.path.join(SRC, "FRANCE Overflight Diplomatic Form_sanitized.xlsx")
    wb = openpyxl.load_workbook(src)
    ws = wb["EU DIC Form"]
    cells = {"M21":"{aircraft}","L22":"{reg}","L24":"{callsign}","L25":"{crew_n}","L26":"{captain}",
             "L31":"{date}","L32":"{purpose}","L33":"{dep}","L34":"{dest}","L38":"{pax}"}
    for coord, ph in cells.items():
        ws[coord] = ph
    outdir = os.path.join(ROOT, "Master Forms", "France"); os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, "france.xlsx"); wb.save(out)
    scrub_docprops(out)
    print(f"  France       xlsx cells set={len(cells)}  -> {os.path.relpath(out, ROOT)}")
    return "france", out

def build_pdfs():
    """PDF templates for the pdf-lib overlay engine (runtime white-out + coordinate overlay).
    - NORMALISE (pikepdf/qpdf): PDFs pdf-lib can already load (NZ, Bahrain).
    - RASTERISE+WRAP: PDFs pdf-lib can't load (Jordan = corrupt source) -> render to PNG, wrap
      in a clean image-PDF at the right page size so pdf-lib loads it; overlay coords still in points.
    Coordinate maps live in index.html (PDFMAP)."""
    import pikepdf, warnings, glob; warnings.simplefilter("ignore")
    from reportlab.pdfgen import canvas
    out = []
    A4 = (595.44, 841.68)
    NORM = {"nz": ("NZ", "NZ diplomatic-aircraft-clearance-form_sanitized_sanitized_sanitized.pdf"),
            "bahrain": ("Bahrain", "BAHRAIN_sanitized.pdf")}
    for key, (folder, fn) in NORM.items():
        pdf = pikepdf.open(os.path.join(SRC, fn))
        outdir = os.path.join(ROOT, "Master Forms", folder); os.makedirs(outdir, exist_ok=True)
        p = os.path.join(outdir, key + ".pdf")
        pdf.save(p, object_stream_mode=pikepdf.ObjectStreamMode.disable, force_version="1.5")
        if key == "nz":
            # The NZ master is a real filed request: page 1 carries a named officer's contact
            # details, a sample captain, callsign, itinerary and routes. Runtime white-out only
            # HIDES those (still copy-pasteable), so bake the redaction in and bundle that
            # instead of the normalised master. See tools/flatten_nz.py.
            import flatten_nz
            flatten_nz.main()
            p = os.path.join(outdir, "nz_flat.pdf")
        out.append((key, p)); print(f"  {key:10} pdf normalized -> {os.path.relpath(p, ROOT)} ({os.path.getsize(p)}b)")
    RAST = {"jordan": ("Jordan", "JORDAN TEMPLATE_sanitized.pdf", A4)}
    for key, (folder, fn, size) in RAST.items():
        os.system(f'pdftoppm -png -r 150 -f 1 -l 1 "{os.path.join(SRC, fn)}" /tmp/{key}_raster >/dev/null 2>&1')
        png = sorted(glob.glob(f"/tmp/{key}_raster*.png"))[0]
        outdir = os.path.join(ROOT, "Master Forms", folder); os.makedirs(outdir, exist_ok=True)
        p = os.path.join(outdir, key + ".pdf")
        c = canvas.Canvas(p, pagesize=size); c.drawImage(png, 0, 0, width=size[0], height=size[1]); c.save()
        out.append((key, p)); print(f"  {key:10} pdf rasterized+wrapped -> {os.path.relpath(p, ROOT)} ({os.path.getsize(p)}b)")
    return out

def build_xfa():
    """Australia & Canada ship from their governments as Adobe XFA (LiveCycle) forms that
    render ONLY in Adobe Reader — every other viewer shows "Please wait…", and current
    Adobe Reader (2024/25) itself now fails them with "STOP! Javascript is not enabled".
    So we render them offline to normal multi-page PDFs with tools/xfa_flatten/ (pdf.js via
    headless Chrome — patch_tpl.py reveals the script-gated FormPage/hides the JS-STOP +
    dev subMetaData; build_artifacts.py renders + extracts marker coords):
      Australia -> australia_flat.pdf (6-page A4 blank); filled at RUNTIME by index.html
                   PDFMAP.australia at coords from tools/xfa_flatten/aus_coords.json.
      Canada    -> canada_flat.pdf   (2-page flat blank; dynamic form, complete in any viewer).
    The XFA *masters* Master Forms/{Australia/australia.pdf, Canada/canada.pdf} are the
    flatten inputs. Here we just bundle the pre-flattened blanks under keys australia/canada;
    re-run tools/xfa_flatten/build_artifacts.py if the source XFA changes.
    (Pre-2026-08 this bundled the raw XFA + a datasets marker-fill; replaced because no
    non-Adobe viewer — and not even the installed Adobe Reader — will render XFA.)"""
    out = []
    for key, rel in [("australia", os.path.join("Master Forms", "Australia", "australia_flat.pdf")),
                     ("canada",    os.path.join("Master Forms", "Canada", "canada_flat.pdf"))]:
        fp = os.path.join(ROOT, rel)
        if os.path.exists(fp):
            out.append((key, fp)); print(f"  {key:10} flat blank (pdf.js flatten) -> {rel} ({os.path.getsize(fp)}b)")
        else:
            print(f"  !! {key}: {rel} MISSING — run tools/xfa_flatten/build_artifacts.py first")
    return out

if __name__ == "__main__":
    print("Building templates:")
    built = [build(c) for c in CONFIGS]
    built.append(build_france())
    built += build_pdfs()
    built += build_xfa()
    keys = {"indonesia": os.path.join(ROOT, "Master Forms", "Indonesia", "indonesia.docx")}
    # Australia's LIVE XFA form, scrubbed and markered by xfa_flatten/patch_tpl.py. Australia
    # will not accept the flattened render, so this ships alongside it and the runtime injects
    # values into its datasets packet, leaving /XFA and /NeedsRendering intact.
    aus_xfa = os.path.join(ROOT, "Master Forms", "Australia", "australia_xfa.pdf")
    if os.path.exists(aus_xfa): keys["australia_xfa"] = aus_xfa
    scrub_docprops(keys["indonesia"])          # bundled directly, so it misses build()'s scrub
    for k, path in built: keys[k] = path
    lines = ["window.TEMPLATES=window.TEMPLATES||{};"]
    for k, path in keys.items():
        lines.append(f'window.TEMPLATES.{k}="{base64.b64encode(open(path,"rb").read()).decode()}";')
    open(os.path.join(ROOT, "lib", "templates_data.js"), "w").write("\n".join(lines) + "\n")
    print(f"\nBundled {len(keys)} templates -> lib/templates_data.js: {list(keys)}")
    stamp_cache_busters()
